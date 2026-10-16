"""从 slides.json 生成可编辑 PPT、文字版、预览和 Notebook 模板。

图和文字使用同一套坐标；PPT 里的文字和流程图均为原生可编辑对象。
预览依赖本机中文字体，是排版参考，不是 PowerPoint 导出的截图。
"""

import json
import math
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from lesson_content import draw_diagram, build_simple_notebook


ROOT = Path(__file__).resolve().parents[1]

BG, INK, TEAL = "F7F5EF", "132C3E", "087F8C"
MUTED, GOLD, LIGHT, BLUE = "5C6F7B", "DAA35B", "E6EFEE", "DEE8F0"
WIDTH, HEIGHT, SCALE = 13.333333, 7.5, 120


def find_font():
    candidates = [os.environ.get("CLASS1_PREVIEW_FONT", ""),
                  "/System/Library/Fonts/PingFang.ttc",
                  "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
                  "C:/Windows/Fonts/msyh.ttc",
                  "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise RuntimeError("预览需要中文字体。请用 CLASS1_PREVIEW_FONT 指定字体文件路径。")


FONT = None


class Canvas:
    def __init__(self, slide):
        self.slide = slide
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        self.image = Image.new("RGB", (1600, 900), "#" + BG)
        self.draw = ImageDraw.Draw(self.image)

    def rect(self, x, y, w, h, color=LIGHT, rounded=True):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        if rounded:
            shape.adjustments[0] = 0.06
        bounds = (round(x*SCALE), round(y*SCALE), round((x+w)*SCALE), round((y+h)*SCALE))
        if rounded:
            self.draw.rounded_rectangle(bounds, radius=10, fill="#"+color)
        else:
            self.draw.rectangle(bounds, fill="#"+color)

    def text(self, x, y, w, h, content, size=20, color=INK, bold=False, center=False):
        font = ImageFont.truetype(FONT, round(size*SCALE/72))
        lines = []
        for original in str(content).split("\n"):
            current = ""
            for char in original:
                if self.draw.textlength(current + char, font=font) > (w*SCALE-6):
                    lines.append(current)
                    current = char
                else:
                    current += char
            lines.append(current)
        line_px = size*SCALE/72*1.3
        if len(lines)*line_px > h*SCALE + 2:
            raise ValueError(f"文字超出预留高度：{content!r}")
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = False  # 使用已经按字体测量好的换行。
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        from pptx.enum.text import PP_ALIGN
        for i, line in enumerate(lines):
            paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.line_spacing = Pt(size*1.3)
            paragraph.space_before = paragraph.space_after = Pt(0)
            if center:
                paragraph.alignment = PP_ALIGN.CENTER
            px = x*SCALE
            if center:
                px += (w*SCALE - self.draw.textlength(line, font=font))/2
            self.draw.text((round(px), round(y*SCALE+i*line_px)), line,
                           font=font, fill="#"+color)

    def line(self, x1, y1, x2, y2, color=TEAL, arrow=False):
        line = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = RGBColor.from_string(color)
        line.line.width = Pt(1.7)
        self.draw.line((x1*SCALE, y1*SCALE, x2*SCALE, y2*SCALE), fill="#"+color, width=3)
        if arrow:
            angle = math.atan2(y2-y1, x2-x1)
            points = [(x2*SCALE, y2*SCALE)]
            for delta in (-0.55, 0.55):
                points.append(((x2-0.09*math.cos(angle+delta))*SCALE,
                               (y2-0.09*math.sin(angle+delta))*SCALE))
            self.draw.polygon(points, fill="#"+color)
            tri = self.slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                Inches(x2-0.035), Inches(y2-0.045), Inches(0.07), Inches(0.09))
            tri.rotation = math.degrees(angle) + 90
            tri.fill.solid()
            tri.fill.fore_color.rgb = RGBColor.from_string(color)
            tri.line.fill.background()

    def node(self, x, y, w, h, text, color=LIGHT, size=17, text_color=INK):
        self.rect(x, y, w, h, color)
        self.text(x+0.08, y+0.09, w-0.16, h-0.10, text, size, text_color, center=True)


def build_slides(slides):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(WIDTH), Inches(HEIGHT)
    presentation.core_properties.title = "Class 1 — Transformer：从一句话到下一个 token"
    presentation.core_properties.subject = "60 分钟入门教学，简明版：文字、加权混合、下一 token"
    presentation.core_properties.author = "Transformer Newbie"
    previews, markdown, minute = [], ["# Class 1 PPT 文字版\n", "由 `slides.json` 自动生成；备注也包含在 PPT 中。\n"], 0
    for number,item in enumerate(slides,1):
        start,end=minute,minute+item["minutes"]
        minute=end
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        c=Canvas(slide)
        c.rect(0,0,.13,HEIGHT,TEAL,False)
        c.text(.65,.35,8,.35,item["section"],13,TEAL,bold=True)
        c.text(10.1,.35,2.58,.35,f"{start:02d}–{end:02d} min",13,MUTED,center=True)
        c.text(.65,.92,12.05,.82,item["title"],30,bold=True)
        c.rect(7.02,1.86,5.65,4.06,"FFFFFF")
        code=item.get("code")
        for i,bullet in enumerate(item["bullets"]):
            y=2.02+i*(.87 if code else 1.05)
            c.node(.65,y,.38,.38,str(i+1),GOLD,13)
            c.text(1.2,y+.02,5.45,.93,bullet,19)
        if code:
            c.rect(.65,4.95,6.03,.94,INK)
            c.text(.83,5.08,5.66,.72,code,14,"FFFFFF")
        draw_diagram(c,item["kind"])
        c.rect(.65,6.12,12.02,.69,LIGHT)
        c.text(.87,6.28,11.57,.45,item["callout"],18,TEAL,bold=True)
        c.text(.65,7.07,9.2,.25,"TRANSFORMER NEWBIE  /  CLASS 1  /  先预测 · 再运行 · 再解释",10,MUTED)
        c.text(11.73,7.07,.92,.25,f"{number:02d} / {len(slides):02d}",10,MUTED,center=True)
        slide.notes_slide.notes_text_frame.text=f"第 {number} 页｜{start}–{end} 分钟\n{item['notes']}"
        previews.append(c.image)
        markdown += [f"\n## {number:02d}. {item['title']}（{start}–{end} 分钟）\n",
                     "\n".join(f"- {b}" for b in item["bullets"])+"\n"]
        if code:
            markdown.append("\n```python\n"+code+"\n```\n")
        markdown += ["\n> "+item["callout"]+"\n", "\n**讲者备注**："+item["notes"]+"\n"]
    if minute != 60:
        raise ValueError(f"课程时长应为 60 分钟，实际为 {minute}")
    presentation.save(ROOT/"class1_simple.pptx")
    (ROOT/"slides.md").write_text("\n".join(markdown).rstrip()+"\n",encoding="utf-8")
    assets=ROOT/"assets"
    assets.mkdir(exist_ok=True)
    sheet=Image.new("RGB",(1440,math.ceil(len(previews)/3)*290),"#D9E0E4")
    for i,preview in enumerate(previews):
        thumbnail=preview.copy()
        thumbnail.thumbnail((470,265))
        sheet.paste(thumbnail,((i%3)*480+5,(i//3)*290+5))
    sheet.save(assets/"slides_preview.jpg",quality=90)
    # 单页预览供检查和复用，不将整组重复图片提交到仓库。
    preview_dir=ROOT/"outputs"/"slide_previews"
    preview_dir.mkdir(parents=True,exist_ok=True)
    for i,slide_image in enumerate(previews,1):
        slide_image.save(preview_dir/f"slide_{i:02d}.png")


def main():
    global FONT
    FONT=find_font()
    slides=json.loads((ROOT/"slides.json").read_text(encoding="utf-8"))
    build_slides(slides)
    build_simple_notebook(ROOT)
    print(f"已生成 {len(slides)} 页 / 60 分钟的 PPT、文字版、预览与 Notebook 模板。")


if __name__ == "__main__":
    main()
