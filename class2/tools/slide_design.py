"""第二课投影版：大公式、可读代码和解释机制的原生图形。"""
from PIL import ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from formulas import place_formula

INK, MUTED, BLUE, TEAL = "14263D", "475569", "2458C6", "087F8C"
BG, WHITE, PALE, LINE = "F4F7FB", "FFFFFF", "E8F0FC", "CBD5E1"
GREEN, RED, DARK = "087568", "B6453B", "14263D"

# 标题概括问题，正文展示定义、解释和计算过程。
DESIGN = {
    "opening": ("文本怎样成为模型输入？", "从真实文本出发，追踪 ID、可训练参数和位置三个接口。", "今天沿同一条链路，解释规则、参数和顺序。"),
    "budget": ("分得更细，还是词表更大？", "词表大小 V 与序列长度 T，会影响不同部分的计算成本。", "选择切分粒度，也是在选择参数与计算的开销。"),
    "bpe_train": ("BPE：先统计，再合并", "每轮重新统计当前相邻片段，选择最高频的一对。", "频次来自语料；合并后要重新统计。"),
    "rank": ("BPE 编码：谁的 rank 更小？", "规则已经冻结；每次在当前候选中选择优先级最高的一对。", "merge rank 决定合并顺序，token ID 决定查表地址。"),
    "hf_pipeline": ("把 BPE 接到真实文本上", "ByteLevel 处理字节与边界；BPE 处理合并规则。", "先训练并保存配置，再加载同一配置进行编码。"),
    "artifact": ("ID 必须与参数表配套", "同样大小的词表，也可能把同一片段映射到不同的地址。", "更换 tokenizer 时，必须核对 ID 与模型参数的对应关系。"),
    "onehot": ("Embedding：一个 ID 选择一行", "查表与 one-hot 乘法前向等价；实际实现直接选行。", "ID 是地址；可训练的表示存放在 E 的每一行。"),
    "lookup": ("不同长度，怎样组成 batch？", "先添加起止标记，再补齐；pad 只是占位。", "padding_idx、注意力 mask 和损失忽略，需要分别处理。"),
    "grad": ("读了同一行，梯度就要相加", "共享参数接收所有读取位置的贡献；这里先看 lookup 路径。", "反向是 scatter-add：累加贡献，而不是覆盖旧值。"),
    "chain": ("预测误差怎样传回向量？", "单位置交叉熵：先得到类别误差，再乘回输出权重。", "先算每次出现的向量梯度，再按 ID 汇总到参数表。"),
    "training": ("让 Embedding 真正学一次", "57 对相邻 token；Embedding → Linear → CrossEntropy。", "这是只看当前 token 的 bigram；没有 Attention。"),
    "tie": ("一张表，两条梯度路径", "权重绑定：输入查表和输出分类器共享同一个 E。", "绑定后，没作为输入的行也可能收到输出梯度。"),
    "permutation": ("没有位置，重排会怎样？", "假设：没有位置机制，也没有因果 mask。", "等变表示“输出跟着重排”，不是“输出完全不变”。"),
    "absolute": ("把位置向量加进输入", "内容与位置使用同样的维度，相加后形状保持不变。", "位置是 token 序号，不是原文的字符 offset。"),
    "frequencies": ("位置 × 频率，得到角度", "不同维度对使用不同频率，再交错存放 sin 与 cos。", "positions → angles → sin / cos：每一步都有明确形状。"),
    "rotation": ("RoPE：相对位移从哪里来？", "q、k 分别按绝对位置旋转；点积里出现位置之差。", "转置是反向旋转；角度相加后，剩下 n−m。"),
    "vector_rope": ("把二维旋转写成两行代码", "相邻维度组成一对；同一公式广播到 batch 与 head。", "保持向量范数；点积仍会随相对位置改变。"),
    "mask": ("位置关系与可读范围", "RoPE 改变分数；mask 决定哪些 key 可以被读取。", "有了位置机制，仍要屏蔽未来 key 与 padding key。"),
    "pipeline": ("沿真实输入，把形状接起来", "本地 tokenizer → Embedding → 多头 Q/K → RoPE + mask。", "这是随机参数的接口实验，尚未训练完整 Transformer。"),
    "quiz": ("用三个问题，检查是否理解", "解释一次编码、一次反传、一次位置计算。", "课后：改词表、查梯度、验证旋转，并保留实验结果。"),
}


def canvas_class(base, font_path):
    class ProjectionCanvas(base):
        def __init__(self, slide):
            super().__init__(slide)
            self.rect(0, 0, 13.333333, 7.5, BG, False)
            self.text_metrics = []
            self.font_path = font_path

        def text(self, x, y, w, h, content, size=22, color=INK, bold=False,
                 center=False, mono=False):
            name = "Menlo" if mono else "Arial Unicode MS"
            path = "/System/Library/Fonts/Menlo.ttc" if mono else font_path
            try:
                font = ImageFont.truetype(path, round(size*120/72))
            except OSError:
                font = ImageFont.truetype("DejaVuSansMono.ttf" if mono else font_path, round(size*120/72))
            lines = []
            for original in str(content).split("\n"):
                current = ""
                for char in original:
                    if self.draw.textlength(current+char, font=font) > w*120-5:
                        if mono:
                            raise ValueError(f"代码必须显式换行：{original}")
                        lines.append(current)
                        current = char
                    else:
                        current += char
                lines.append(current)
            step = size*120/72*1.22
            if len(lines) == 1:
                h = max(h, size/72*1.22)
            if len(lines)*step > h*120+2:
                raise ValueError(f"文本超出高度（{size} pt）：{content!r}")
            box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.font.name = name
                p.font.size = Pt(size)
                p.font.bold = bold
                p.font.color.rgb = RGBColor.from_string(color)
                p.line_spacing = Pt(size*1.22)
                p.space_before = p.space_after = Pt(0)
                if center:
                    p.alignment = PP_ALIGN.CENTER
                left = x*120+(w*120-self.draw.textlength(line, font=font))/2 if center else x*120
                self.draw.text((round(left), round(y*120+i*step)), line,
                               font=font, fill="#"+color, stroke_width=1 if bold else 0)
            self.text_metrics.append({"text": str(content), "size": size, "mono": mono,
                                      "lines": len(lines), "box": [x,y,w,h]})

    return ProjectionCanvas


def token(c, x, y, text, w=.8, fill=PALE, color=BLUE, size=24):
    c.rect(x,y,w,.58,fill)
    c.text(x+.06,y+.07,w-.12,.48,text,size,color,center=True)


def header(c, item, number, start, end, total):
    kind = item["kind"]
    title, subtitle, takeaway = DESIGN[kind]
    active = 0 if number <= 6 else 1 if number <= 12 else 2
    for i, text in enumerate(("01  分词", "02  向量", "03  位置")):
        x = .60+i*1.52
        c.rect(x,.27,1.34,.34,BLUE if i == active else "E6ECF4")
        c.text(x,.31,1.34,.26,text,14,WHITE if i == active else MUTED,center=True)
    c.text(10.0,.30,2.7,.28,f"{start:02d}–{end:02d} min  /  60",14,MUTED,center=True)
    c.text(.60,.84,12.10,.66,title,34,INK,bold=True)
    c.text(.62,1.54,12.0,.39,subtitle,20,MUTED)
    c.rect(.60,6.70,12.10,.47,PALE)
    c.text(.79,6.78,11.7,.36,takeaway,19,BLUE,bold=True)
    c.text(.62,7.27,8.8,.21,"TRANSFORMER NEWBIE  ·  CLASS 2",11,MUTED)
    c.text(11.5,7.25,1.2,.24,f"{number:02d} / {total:02d}",12,MUTED,center=True)


def reference(c, link):
    if not link:
        return
    experiment = f"实验 {link['experiment']:02d}" if link["experiment"] else "课后扩展"
    file = link["file"].replace("../tools/", "tools/")
    text=f"{experiment}   ·   {file} / {link['function']}"
    if link.get("worked_function"):
        text=f"{experiment}   ·   {file}   ·   算例 worked_examples.py/{link['worked_function']}"
    c.text(.65,6.36,12,.27,text,14,MUTED)
