"""制作第二课 PPT、讲义文字版和 Notebook。"""
import importlib.util
import json
import math
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from formulas import write_alignment
from lesson_alignment import LINKS, experiment_label
from slide_design import DESIGN, canvas_class, header, reference, save_vector_pdf, validate_layout
from detail_slides import draw_detail

ROOT = Path(__file__).resolve().parents[1]
base_tools = ROOT.parent / "class1/tools"
sys.path.append(str(base_tools))
spec = importlib.util.spec_from_file_location("course_drawing", base_tools / "build_materials.py")
drawing = importlib.util.module_from_spec(spec)
spec.loader.exec_module(drawing)


def build_slides():
    items = json.loads((ROOT / "slides.json").read_text(encoding="utf-8"))
    ppt = Presentation()
    ppt.slide_width, ppt.slide_height = Inches(13.333333), Inches(7.5)
    ppt.core_properties.title = "第二课：Tokenizer、Embedding 与位置机制"
    ppt.core_properties.subject = "Transformer 输入表示，60 分钟"
    ppt.core_properties.author = "Transformer Newbie"
    previews, text, minute = [], ["# 第二课：Tokenizer、Embedding 与位置机制\n"], 0
    Canvas = canvas_class(drawing.Canvas)
    canvases = []
    report = json.loads((ROOT / "assets/training_report.json").read_text(encoding="utf-8"))
    sys.path.insert(0,str(ROOT / "code"))
    from worked_examples import build_report
    worked = build_report()
    layout_report = []
    for number, item in enumerate(items, 1):
        start, end = minute, minute + item["minutes"]
        minute = end
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        c = Canvas(slide)
        header(c,item,number,start,end,len(items))
        link = LINKS.get(item["kind"])
        draw_detail(c,item,link,ROOT,report,worked)
        reference(c,link)
        title, subtitle, takeaway = DESIGN[item["kind"]]
        notes = (f"第 {number} 页｜{start}–{end} 分钟\n{item['title']}\n"
                 +"\n".join(item["bullets"])+f"\n{item['callout']}\n{item['notes']}")
        notes += "\n页面解释：\n"+"\n".join(item.get("explanation",[]))
        notes += "\n数值算例：code/worked_examples.py；结果：assets/worked_examples.json。"
        if link:
            file = link["file"].replace("../tools/", "tools/")
            notes += f"\n对应：{experiment_label(link)}，{file}:{link['function']}\n观察：{link['result']}"
            if item["kind"] == "absolute":
                notes += "\n投影代码出自 position_mechanics.py:demo；位置函数来自 positions.py:sinusoidal。"
        slide.notes_slide.notes_text_frame.text = notes
        previews.append(c.image)
        canvases.append(c)
        layout_report.append({"slide":number,"title":title,"text":c.text_metrics,"ink":c.ink_boxes})
        text += [f"\n## {number:02d}. {title}（{start}–{end} 分钟）\n",
                 subtitle, "\n".join("- "+b for b in item.get("explanation",item["bullets"])), "\n> "+takeaway,
                 "\n**授课提示**："+item["notes"]]
        if link:
            text += ["\n**对应实验**："+experiment_label(link)+"；`"+file+":"+link["function"]+"`。"]
            if link.get("visual_code"):
                text += ["\n**箭头图解**：每格30°；位置1/3和3/5的点积均为0.5。",
                         "\n```python\n"+"\n".join(link["visual_code"])+"\n```"]
            for formula in link.get("formulas", []):
                text += ["\n$$\n"+formula+"\n$$"]
            if link.get("symbols"):
                text += ["\n"+link["symbols"]+"。"]
            if link.get("code"):
                text += ["\n```python\n"+"\n".join(link["code"])+"\n```"]
            text += ["\n**运行观察**："+link["result"]+"。"]
    assert minute == 60
    validation = validate_layout(canvases)
    ppt.save(ROOT / "class2_tokens_embeddings_positions.pptx")
    save_vector_pdf(canvases, ROOT / "class2_tokens_embeddings_positions.pdf")
    (ROOT / "slides.md").write_text("\n".join(text).rstrip()+"\n", encoding="utf-8")
    write_alignment(ROOT, items)
    (ROOT / "assets").mkdir(exist_ok=True)
    sheet = Image.new("RGB", (1920, math.ceil(len(items)/3)*382), "#D9E0E4")
    output = ROOT / "outputs/slide_previews"
    output.mkdir(parents=True, exist_ok=True)
    for i, preview in enumerate(previews):
        preview.save(output / f"slide_{i+1:02d}.png")
        thumbnail = preview.copy()
        thumbnail.thumbnail((628, 354))
        sheet.paste(thumbnail, ((i%3)*640+6, (i//3)*382+6))
    sheet.save(ROOT / "assets/slides_preview.jpg", quality=90)
    (output / "layout_report.json").write_text(json.dumps(layout_report, ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    (output / "validation.json").write_text(json.dumps(validation,indent=2)+"\n")
    # 独立示例供浏览，无需先打开整个 PPT。
    for name, index in (("tokenizer",2),("opening",0),("scope",1),("bpe_training",2),
                        ("bpe_encoding",3),("embedding",8),("rope",15)):
        previews[index].save(ROOT / f"assets/preview_{name}.png")


def build_notebooks():
    from notebooks import build
    build(ROOT)


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--slides-only", action="store_true", help="仅重建 PPT 与预览，保留 Notebook 的执行结果")
    args = parser.parse_args()
    build_slides()
    if not args.slides_only:
        build_notebooks()
    print("第二课课堂阅读版：20 页、60 分钟、完整算例与真实代码。")
